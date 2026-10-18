import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os

TEMPLATE_PATH = r'C:/Users/osmae/.gemini/antigravity/brain/5cfe47df-e6da-43e1-82df-44db0f1814d6/.user_uploaded/media_1792319972213_50052aec.pptx'
OUTPUT_PATH = r'c:\Users\osmae\.gemini\antigravity\scratch\HR-Digital\presentation_team_47.pptx'
PIC1_PATH = r'c:\Users\osmae\.gemini\antigravity\scratch\HR-Digital\scripts\slide3_pic_1.png'
PIC2_PATH = r'c:\Users\osmae\.gemini\antigravity\scratch\HR-Digital\scripts\slide3_pic_2.png'

prs = pptx.Presentation(TEMPLATE_PATH)

# ==========================================
# SLIDE 1: ТИТУЛЬНЫЙ СЛАЙД
# ==========================================
# По требованию пользователя:
# "первый слайд просто заполни как есть, не меняя блоки местами и с картинками пропиши туда наши данные и тд. номер команды 47"
slide1 = prs.slides[0]

# Shape 0: Placeholder "название кейса" (pos: 311708, 744575, 8520600, 2052600)
sh_case = slide1.shapes[0]
tf_case = sh_case.text_frame
tf_case.word_wrap = True
tf_case.clear()

p_title = tf_case.paragraphs[0]
p_title.alignment = PP_ALIGN.CENTER
p_title.space_after = Pt(12)
run_title = p_title.add_run()
run_title.text = "Digital HR: Корпоративная платформа"
run_title.font.name = "Arial"
run_title.font.size = Pt(36)
run_title.font.bold = True
run_title.font.color.rgb = RGBColor(16, 24, 40) # #101828 Deep Navy

p_sub = tf_case.add_paragraph()
p_sub.alignment = PP_ALIGN.CENTER
run_sub = p_sub.add_run()
run_sub.text = "Автоматизация кадровых процессов и учет рабочего времени для Jana Post"
run_sub.font.name = "Arial"
run_sub.font.size = Pt(18)
run_sub.font.bold = False
run_sub.font.color.rgb = RGBColor(71, 84, 103) # Slate gray

# Shape 1: Placeholder "номер команды" (pos: 311700, 2834125, 8520600, 792600)
sh_team = slide1.shapes[1]
tf_team = sh_team.text_frame
tf_team.word_wrap = True
tf_team.clear()

p_team = tf_team.paragraphs[0]
p_team.alignment = PP_ALIGN.CENTER
run_team = p_team.add_run()
run_team.text = "Команда № 47"
run_team.font.name = "Arial"
run_team.font.size = Pt(28)
run_team.font.bold = True
run_team.font.color.rgb = RGBColor(37, 99, 235) # #2563EB Corporate Blue

# Shape 2: Picture (IT-HUB QARAGANDY / Терриконовая долина) - остается нетронутой!
# Shape 3: Text Box "Караганда 2026 г." - остается нетронутой!


# ==========================================
# SLIDE 2: КРАТКОЕ ОПИСАНИЕ ПРОДЕЛАННОЙ РАБОТЫ
# ==========================================
slide2 = prs.slides[1]

# Title shape: Shape 0
sh_title2 = slide2.shapes[0]
tf_title2 = sh_title2.text_frame
tf_title2.clear()
p_t2 = tf_title2.paragraphs[0]
p_t2.alignment = PP_ALIGN.LEFT
r_t2 = p_t2.add_run()
r_t2.text = "Краткое описание проделанной работы"
r_t2.font.name = "Arial"
r_t2.font.size = Pt(28)
r_t2.font.bold = True
r_t2.font.color.rgb = RGBColor(16, 24, 40)

# Body shape: Shape 1
sh_body2 = slide2.shapes[1]
tf_body2 = sh_body2.text_frame
tf_body2.word_wrap = True
tf_body2.clear()

# 1. О проекте
p_desc = tf_body2.paragraphs[0]
p_desc.space_after = Pt(8)
r_d1 = p_desc.add_run()
r_d1.text = "📌 О проекте: "
r_d1.font.name = "Arial"
r_d1.font.size = Pt(13)
r_d1.font.bold = True
r_d1.font.color.rgb = RGBColor(16, 24, 40)
r_d2 = p_desc.add_run()
r_d2.text = "Разработана enterprise-платформа Digital HR для управления кадровым циклом Jana Post в соответствии с Трудовым кодексом РК (UTC+5, квота отпусков 24 дня, формы Т-2, Т-6, Т-13)."
r_d2.font.name = "Arial"
r_d2.font.size = Pt(12)
r_d2.font.color.rgb = RGBColor(51, 65, 85)

# 2. Какие ресурсы были использованы?
p_q1 = tf_body2.add_paragraph()
p_q1.space_before = Pt(6)
p_q1.space_after = Pt(4)
r_q1 = p_q1.add_run()
r_q1.text = "🛠 Какие ресурсы были использованы?"
r_q1.font.name = "Arial"
r_q1.font.size = Pt(14)
r_q1.font.bold = True
r_q1.font.color.rgb = RGBColor(37, 99, 235)

bullets_q1 = [
    ("• Frontend & UI: ", "React 19, TypeScript, Vite, Tailwind CSS v4, Lucide React, TanStack Query v5, Recharts (аналитика смен и отпусков)."),
    ("• Backend & Данные: ", "Node.js, Express, REST API v1, Prisma ORM, PostgreSQL / SQLite (автономный режим), JWT-авторизация с ролевым доступом (RBAC)."),
    ("• AI & База знаний: ", "Кадровый RAG-ассистент «Jana» со встроенной верифицированной базой регламентов РК и безопасным отказом.")
]

for bold_prefix, text_content in bullets_q1:
    p = tf_body2.add_paragraph()
    p.space_after = Pt(3)
    rb = p.add_run()
    rb.text = bold_prefix
    rb.font.name = "Arial"
    rb.font.size = Pt(11.5)
    rb.font.bold = True
    rb.font.color.rgb = RGBColor(30, 41, 59)
    rt = p.add_run()
    rt.text = text_content
    rt.font.name = "Arial"
    rt.font.size = Pt(11.5)
    rt.font.color.rgb = RGBColor(71, 84, 103)

# 3. Какие результаты получили
p_q2 = tf_body2.add_paragraph()
p_q2.space_before = Pt(6)
p_q2.space_after = Pt(4)
r_q2 = p_q2.add_run()
r_q2.text = "🎯 Какие результаты получили"
r_q2.font.name = "Arial"
r_q2.font.size = Pt(14)
r_q2.font.bold = True
r_q2.font.color.rgb = RGBColor(16, 185, 129) # Emerald Green

bullets_q2 = [
    ("• 3 сквозных бизнес-сценария: ", "Личный кабинет сотрудника (ESS) с трекером прихода/ухода, очередь согласования руководителя команды (SLA 2 ч) и модуль кадрового администрирования."),
    ("• Безопасность и изоляция: ", "Строгая серверная проверка ролей на каждом эндпоинте; маскирование ИИН, защита от подмены ID сотрудников и полный аудит-лог."),
    ("• Интерактивный прототип: ", "100% готовность к демонстрации на реальных seed-данных в автономном режиме без внешней сети.")
]

for bold_prefix, text_content in bullets_q2:
    p = tf_body2.add_paragraph()
    p.space_after = Pt(3)
    rb = p.add_run()
    rb.text = bold_prefix
    rb.font.name = "Arial"
    rb.font.size = Pt(11.5)
    rb.font.bold = True
    rb.font.color.rgb = RGBColor(30, 41, 59)
    rt = p.add_run()
    rt.text = text_content
    rt.font.name = "Arial"
    rt.font.size = Pt(11.5)
    rt.font.color.rgb = RGBColor(71, 84, 103)


# ==========================================
# SLIDE 3: ДЕМОНСТРАЦИЯ ПРОТОТИПА
# ==========================================
slide3 = prs.slides[2]

# Title shape: Shape 0
sh_title3 = slide3.shapes[0]
tf_title3 = sh_title3.text_frame
tf_title3.clear()
p_t3 = tf_title3.paragraphs[0]
p_t3.alignment = PP_ALIGN.LEFT
r_t3 = p_t3.add_run()
r_t3.text = "Демонстрация прототипа"
r_t3.font.name = "Arial"
r_t3.font.size = Pt(28)
r_t3.font.bold = True
r_t3.font.color.rgb = RGBColor(16, 24, 40)

# Body shape: Shape 1 (Set width to ~45% of slide width to leave room for screenshots)
sh_body3 = slide3.shapes[1]
sh_body3.left = 311700
sh_body3.top = 1152475
sh_body3.width = 4200000
sh_body3.height = 3750000

tf_body3 = sh_body3.text_frame
tf_body3.word_wrap = True
tf_body3.clear()

p_scen_title = tf_body3.paragraphs[0]
p_scen_title.space_after = Pt(6)
r_st = p_scen_title.add_run()
r_st.text = "Ключевые сценарии демонстрации:"
r_st.font.name = "Arial"
r_st.font.size = Pt(14)
r_st.font.bold = True
r_st.font.color.rgb = RGBColor(37, 99, 235)

scenarios = [
    ("1. Кабинет сотрудника (ESS):", [
        "• Расчет баланса отпусков (ст. 88 ТК РК — 24 дня)",
        "• Подача заявления на отпуск и запрос аванса",
        "• Фиксация рабочего времени Clock In / Out (UTC+5)"
    ]),
    ("2. Кабинет руководителя отдела:", [
        "• Очередь заявлений сотрудников IT-подразделения",
        "• Резолюция в 1 клик: Одобрить / Отклонить с аудитом"
    ]),
    ("3. AI-ассистент «Jana»:", [
        "• RAG-ответы со ссылками на первоисточники ТК РК",
        "• Безопасный отказ при непрофильных вопросах"
    ])
]

for section_title, items in scenarios:
    p_sec = tf_body3.add_paragraph()
    p_sec.space_before = Pt(5)
    p_sec.space_after = Pt(2)
    r_sec = p_sec.add_run()
    r_sec.text = section_title
    r_sec.font.name = "Arial"
    r_sec.font.size = Pt(12)
    r_sec.font.bold = True
    r_sec.font.color.rgb = RGBColor(16, 24, 40)
    
    for item in items:
        p_item = tf_body3.add_paragraph()
        p_item.space_after = Pt(1)
        r_item = p_item.add_run()
        r_item.text = item
        r_item.font.name = "Arial"
        r_item.font.size = Pt(10.5)
        r_item.font.color.rgb = RGBColor(71, 84, 103)

p_acc = tf_body3.add_paragraph()
p_acc.space_before = Pt(8)
r_acc = p_acc.add_run()
r_acc.text = "👤 Тестовые роли:\nemployee@janapost.kz / manager@janapost.kz"
r_acc.font.name = "Arial"
r_acc.font.size = Pt(10)
r_acc.font.bold = True
r_acc.font.color.rgb = RGBColor(100, 116, 139)

# Удалим старые картинки на 3 слайде, если они были
shapes_to_remove = []
for sh in slide3.shapes:
    if sh.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
        shapes_to_remove.append(sh)

for sh in shapes_to_remove:
    sp = sh._element
    sp.getparent().remove(sp)

# Добавляем скриншоты прототипа в правую часть слайда
if os.path.exists(PIC1_PATH) and os.path.exists(PIC2_PATH):
    img1_left = 4700000
    img1_top = 1150000
    img1_width = 4150000
    img1_height = int(img1_width * (430 / 1160)) # соотношение сторон
    slide3.shapes.add_picture(PIC1_PATH, img1_left, img1_top, width=img1_width, height=img1_height)

    img2_left = 4700000
    img2_top = 2850000
    img2_width = 4150000
    img2_height = int(img2_width * (420 / 1160)) # соотношение сторон
    slide3.shapes.add_picture(PIC2_PATH, img2_left, img2_top, width=img2_width, height=img2_height)

prs.save(OUTPUT_PATH)
print("Successfully generated presentation at:", OUTPUT_PATH)
