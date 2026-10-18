import pptx
import sys
import json
import os

sys.stdout.reconfigure(encoding='utf-8')

uploaded_path = r'C:/Users/osmae/.gemini/antigravity/brain/5cfe47df-e6da-43e1-82df-44db0f1814d6/.user_uploaded/media_1792319972213_50052aec.pptx'
prs = pptx.Presentation(uploaded_path)
print(f"Presentation width: {prs.slide_width}, height: {prs.slide_height}")

for s_idx, slide in enumerate(prs.slides):
    print(f"\n==================== SLIDE {s_idx + 1} ====================")
    for sh_idx, shape in enumerate(slide.shapes):
        print(f"\n--- Shape {sh_idx}: {shape.name} (type: {shape.shape_type}) ---")
        print(f"    Left: {shape.left}, Top: {shape.top}, Width: {shape.width}, Height: {shape.height}")
        if shape.has_text_frame:
            tf = shape.text_frame
            print(f"    Text frame paragraphs: {len(tf.paragraphs)}")
            for p_idx, p in enumerate(tf.paragraphs):
                p_text = p.text
                print(f"      P{p_idx} [{p.alignment}]: {repr(p_text)}")
                for r_idx, r in enumerate(p.runs):
                    print(f"        Run {r_idx}: {repr(r.text)} (font={r.font.name}, size={r.font.size}, bold={r.font.bold}, color={r.font.color.rgb if r.font.color and r.font.color.type == 1 else 'other'})")
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            image = shape.image
            print(f"    Picture filename: {image.filename}, ext: {image.ext}, size: {len(image.blob)} bytes")
